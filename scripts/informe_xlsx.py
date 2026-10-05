"""
Arma informe/Kiwano_red_de_nodos.xlsx a partir de informe/red-kiwano.json.
Uso: python scripts/informe_xlsx.py  (después de exportar-red.ts)
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
datos = json.loads((RAIZ / "informe" / "red-kiwano.json").read_text(encoding="utf-8"))

VERDE = "1B4D3E"
CAFE = "8C7355"
CREMA = "F7F4EF"
TIPO_COLOR = {"sede": "9dbf85", "cliente": "D6E8DF", "proveedor": "E9DFD2", "competencia": "E4E4E4"}

cab = Font(bold=True, color="FFFFFF")
cab_fill = PatternFill("solid", fgColor=VERDE)
titulo = Font(bold=True, size=14, color=VERDE)
fino = Side(style="thin", color="CCCCCC")
borde = Border(top=fino, bottom=fino, left=fino, right=fino)
COP = '"$"#,##0'


def encabezado(ws, fila, columnas):
    for j, texto in enumerate(columnas, 1):
        c = ws.cell(row=fila, column=j, value=texto)
        c.font, c.fill, c.border = cab, cab_fill, borde
        c.alignment = Alignment(wrap_text=True, vertical="center")


def anchos(ws, valores):
    for j, w in enumerate(valores, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


wb = Workbook()

# 1. Base de datos de nodos
ws = wb.active
ws.title = "Nodos"
ws["A1"] = "Kiwano: base de datos de nodos (Bogotá)"
ws["A1"].font = titulo
ws["A2"] = "N00 es la sede de Kiwano. Coordenadas aproximadas por dirección; horario marcado como estimado cuando no es público."
cols = ["ID", "Empresa", "Tipo", "Sector / servicio", "Dirección", "Localidad", "Latitud", "Longitud",
        "Abre", "Cierra", "Horario estimado", "Relevancia para Kiwano", "Fuente"]
encabezado(ws, 4, cols)
for i, n in enumerate(datos["nodos"], 5):
    fila = [n["id"], n["nombre"], n["tipo"].capitalize(), n["sector"], n["direccion"], n["localidad"],
            n["lat"], n["lon"], n["abre"], n["cierra"], "Sí" if n["horarioEstimado"] else "No",
            n["relevancia"], n["fuente"]]
    for j, v in enumerate(fila, 1):
        c = ws.cell(row=i, column=j, value=v)
        c.border = borde
        c.alignment = Alignment(wrap_text=j in (2, 4, 5, 12), vertical="top")
    ws.cell(row=i, column=3).fill = PatternFill("solid", fgColor=TIPO_COLOR[n["tipo"]])
anchos(ws, [6, 30, 12, 26, 32, 16, 10, 10, 7, 7, 10, 44, 40])
ws.freeze_panes = "C5"

resumen = ws.max_row + 2
ws.cell(row=resumen, column=2, value="Resumen por tipo").font = Font(bold=True, color=VERDE)
for k, tipo in enumerate(["cliente", "proveedor", "competencia"], 1):
    ws.cell(row=resumen + k, column=2, value=tipo.capitalize())
    ws.cell(row=resumen + k, column=3, value=sum(1 for n in datos["nodos"] if n["tipo"] == tipo))

# 2. Matriz de distancias
ws = wb.create_sheet("Distancias")
s = datos["supuestos"]
ws["A1"] = "Matriz de distancias viales estimadas (km)"
ws["A1"].font = titulo
ws["A2"] = (f"Distancia = haversine (línea recta sobre la Tierra) × factor de circuidad {s['factorCircuidad']} "
            "para aproximar el recorrido por calles de Bogotá.")
ids = [n["id"] for n in datos["nodos"]]
encabezado(ws, 4, ["km"] + ids)
for i, (n, fila) in enumerate(zip(datos["nodos"], datos["matriz"]), 5):
    c = ws.cell(row=i, column=1, value=f'{n["id"]} {n["corto"]}')
    c.font, c.fill = cab, cab_fill
    for j, d in enumerate(fila, 2):
        cel = ws.cell(row=i, column=j, value=round(d, 2))
        cel.number_format = "0.0"
        cel.border = borde
        cel.alignment = Alignment(horizontal="center")
ult = get_column_letter(len(ids) + 1)
ws.conditional_formatting.add(
    f"B5:{ult}{4 + len(ids)}",
    ColorScaleRule(start_type="min", start_color="D6E8DF", mid_type="percentile", mid_value=50,
                   mid_color="F7F4EF", end_type="max", end_color="C9B79C"),
)
anchos(ws, [22] + [6.5] * len(ids))
ws.freeze_panes = "B5"

# 3. Ventanas de tiempo y ruta
ws = wb.create_sheet("Ruta y ventanas")
ws["A1"] = "Plan de visitas comerciales con ventanas de tiempo"
ws["A1"].font = titulo
ws["A2"] = (f"Heurística del vecino más cercano factible (VRPTW). Velocidad {s['velocidadKmH']} km/h, "
            f"{s['visitaMin']} min por visita, jornada {s['salida']}–{s['finJornada']}. "
            "Si se llega antes de que abra el cliente, se espera.")
encabezado(ws, 4, ["Día", "Orden", "ID", "Cliente", "Ventana", "Llegada", "Inicio visita", "Km desde parada anterior"])
f = 5
for j in datos["plan"]:
    for k, p in enumerate(j["paradas"], 1):
        for col, v in enumerate([j["dia"], k, p["id"], p["nombre"], p["ventana"], p["llegada"], p["inicio"],
                                 round(p["km"], 1)], 1):
            ws.cell(row=f, column=col, value=v).border = borde
        f += 1
    ws.cell(row=f, column=4, value=f"Regreso a la sede {j['regreso']} · total del día").font = Font(bold=True)
    ws.cell(row=f, column=8, value=round(j["km"], 1)).font = Font(bold=True)
    f += 2
anchos(ws, [6, 7, 6, 34, 13, 9, 12, 14])

# 4. Costos
ws = wb.create_sheet("Costos")
co = datos["costos"]
ws["A1"] = f"Estructura de costos mensual de Kiwano (COP, {co['fecha']})"
ws["A1"].font = titulo
ws["A2"] = co["notaCarga"]
encabezado(ws, 4, ["Tipo", "Grupo", "Concepto", "Valor base", "Con carga prestacional", "Fuente"])
f = 5
for x in co["fijosCalculados"]:
    for col, v in enumerate(["Fijo", x["grupo"], x["concepto"], x["valor"], x["total"], x["fuente"]], 1):
        c = ws.cell(row=f, column=col, value=v)
        c.border = borde
        if col in (4, 5):
            c.number_format = COP
    f += 1
ws.cell(row=f, column=3, value="Total costos fijos").font = Font(bold=True)
c = ws.cell(row=f, column=5, value=co["totalFijos"])
c.number_format, c.font = COP, Font(bold=True)
f += 2
precio = next(p["precio"] for p in co["precios"] if p["plan"] == "Agente")
for x in co["variablesCliente"]:
    valor = x["valor"] * precio if x["tipo"] == "pct" else x["valor"]
    det = f' ({x["valor"] * 100:.2f} % del precio)' if x["tipo"] == "pct" else ""
    for col, v in enumerate(["Variable por cliente", "Cliente", x["concepto"] + det, valor, valor, x["fuente"]], 1):
        c = ws.cell(row=f, column=col, value=v)
        c.border = borde
        if col in (4, 5):
            c.number_format = COP
    f += 1
for x in co["variablesCampana"] + [{"concepto": f'{co["transporte"]["concepto"]} · {co["kmMes"]:.0f} km/mes',
                                     "valor": co["transporteMes"], "fuente": co["transporte"]["fuente"]}]:
    for col, v in enumerate(["Variable por campaña", "Promoción", x["concepto"], x["valor"], x["valor"], x["fuente"]], 1):
        c = ws.cell(row=f, column=col, value=v)
        c.border = borde
        if col in (4, 5):
            c.number_format = COP
    f += 1
f += 1
for etiqueta, valor in [
    ("Precio plan Agente (mensual)", precio),
    ("Costo variable por cliente", co["cvAgente"]),
    ("Margen de contribución por cliente", co["margenAgente"]),
    ("Fijos + campaña del mes", co["totalFijos"] + co["totalCampana"]),
    ("Punto de equilibrio (clientes plan Agente)", co["equilibrio"]),
]:
    ws.cell(row=f, column=3, value=etiqueta).font = Font(bold=True)
    c = ws.cell(row=f, column=5, value=valor)
    c.font = Font(bold=True, color=VERDE)
    if "clientes" not in etiqueta:
        c.number_format = COP
    f += 1
anchos(ws, [20, 12, 58, 14, 18, 50])

# 5. Precios y prompt del gráfico
ws = wb.create_sheet("Precios y prompt")
ws["A1"] = "Planes de precio"
ws["A1"].font = titulo
encabezado(ws, 3, ["Plan", "Precio", "Unidad", "Incluye"])
for i, p in enumerate(co["precios"], 4):
    for col, v in enumerate([p["plan"], p["precio"], p["unidad"], p["incluye"]], 1):
        c = ws.cell(row=i, column=col, value=v)
        c.border = borde
        if col == 2:
            c.number_format = COP
r = 5 + len(co["precios"])
ws.cell(row=r, column=1, value="Prompt usado para generar el gráfico de la red con Claude").font = titulo
c = ws.cell(row=r + 1, column=1, value=datos["promptGrafico"])
c.alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 4, end_column=4)
anchos(ws, [14, 14, 12, 70])

# 6. P-medianas
ws = wb.create_sheet("P-medianas")
pm = datos["pmedianas"]
ws["A1"] = "P-medianas: ubicación del punto de atención"
ws["A1"].font = titulo
ws["A2"] = "Minimiza la suma de distancias de cada cliente a su sitio más cercano. Candidatos: los 21 nodos; solución exacta por enumeración."
encabezado(ws, 4, ["Escenario", "Suma de distancias (km)", "Ubicación"])
filas = [("Sede actual (WeWork Calle 93)", pm["totalDesdeSede"], "N00 Kiwano (Chapinero – Parque 93)")]
filas += [(f"p = {s['p']}", s["total"], " + ".join(s["sitios"])) for s in pm["soluciones"]]
for k, (a, b, c) in enumerate(filas, 5):
    for col, v in enumerate([a, round(b, 1), c], 1):
        ws.cell(row=k, column=col, value=v).border = borde
f = 5 + len(filas) + 1
for s in pm["soluciones"]:
    ws.cell(row=f, column=1, value=f"Asignación con p = {s['p']}").font = Font(bold=True, color=VERDE)
    encabezado(ws, f + 1, ["Cliente", "Sitio asignado", "Km"])
    for k, a in enumerate(s["asignacion"], f + 2):
        for col, v in enumerate([a["cliente"], a["sitio"], round(a["km"], 1)], 1):
            ws.cell(row=k, column=col, value=v).border = borde
    f = f + 3 + len(s["asignacion"])
anchos(ws, [30, 24, 60])

# 7. Modelo de transporte (esquina noroeste, costo mínimo, Vogel)
ws = wb.create_sheet("Transporte")
tr = datos["transporte"]
ws["A1"] = "Modelo de transporte: 3 puntos de despacho, 4 clientes"
ws["A1"].font = titulo
ws["A2"] = "Celda = cantidad asignada (costo unitario entre paréntesis). Vogel coincide con el óptimo (verificado con MODI)."
f = 4
for m in tr["metodos"]:
    ws.cell(row=f, column=1, value=m["nombre"]).font = Font(bold=True, color=VERDE)
    ws.cell(row=f, column=7, value=f"Costo total: {m['total']:,}".replace(",", ".")).font = Font(bold=True)
    encabezado(ws, f + 1, [""] + tr["destinos"] + ["Oferta"])
    for i, o in enumerate(tr["origenes"]):
        ws.cell(row=f + 2 + i, column=1, value=o).font = Font(bold=True)
        for j in range(len(tr["destinos"])):
            q = m["x"][i][j]
            c = ws.cell(row=f + 2 + i, column=2 + j, value=f"{q} ({tr['costo'][i][j]})" if q else f"– ({tr['costo'][i][j]})")
            c.border = borde
            c.alignment = Alignment(horizontal="center")
            if q:
                c.fill = PatternFill("solid", fgColor="D6E8DF")
        ws.cell(row=f + 2 + i, column=2 + len(tr["destinos"]), value=tr["oferta"][i])
    r = f + 2 + len(tr["origenes"])
    ws.cell(row=r, column=1, value="Demanda").font = Font(bold=True)
    for j, d in enumerate(tr["demanda"]):
        ws.cell(row=r, column=2 + j, value=d)
    f = r + 3
anchos(ws, [18, 12, 12, 12, 12, 10, 22])

salida = RAIZ / "informe" / "Kiwano_red_de_nodos.xlsx"
wb.save(salida)
print(f"OK: {salida}")
